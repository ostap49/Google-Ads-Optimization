"""generate_report: Markdown + HTML (embedded charts) + optional XLSX."""
from __future__ import annotations

import base64
import html
import re
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

from core import accounts, db
from core.config import get_settings
from core.fmt import pct_delta, r
from gads.queries import performance as Q
from memory.brief import targets
from reports import charts
from reports.i18n import T
from tools.common import account_header, agg, ctx, deltas, m

KPI_ROWS = ("cost", "conversions", "conv_value", "cpa", "roas", "clicks", "impressions", "ctr", "cpc", "conv_rate", "search_is")
LOWER_IS_BETTER = {"cost", "cpa", "cpc"}


def _fmt(v: Any, key: str = "") -> str:
    if v is None:
        return "—"
    if isinstance(v, float):
        s = f"{v:,.2f}" if key not in ("clicks", "impressions") else f"{v:,.0f}"
    elif isinstance(v, int):
        s = f"{v:,}"
    else:
        return str(v)
    return s.replace(",", " ")


def _delta(v: Any) -> str:
    return "—" if v is None else (f"+{v}%" if v > 0 else f"{v}%")


def _md_table(head: list[str], rows: list[list[Any]]) -> str:
    out = ["| " + " | ".join(head) + " |", "|" + "|".join("---" for _ in head) + "|"]
    out += ["| " + " | ".join(str(c).replace("|", "\\|") for c in row) + " |" for row in rows]
    return "\n".join(out)


def _slug(s: str) -> str:
    return re.sub(r"[^\w-]+", "-", s, flags=re.U).strip("-").lower() or "account"


def build_report(customer_id: str, period: str = "LAST_MONTH", lang: str = "uk", xlsx: bool = False,
                 include_next_steps: bool = True, out_dir: str | None = None) -> dict[str, Any]:
    lang = lang if lang in T else "uk"
    t = T[lang]
    g, cid, info, dr = ctx(customer_id, period)
    tg = targets(cid)
    cur_rows = g.search(cid, Q.account_totals(dr))
    from tools.common import is_metrics
    cur = {**agg(cur_rows), **(is_metrics(cur_rows[0]) if cur_rows else {})}
    pdr = dr.previous()
    prev_rows = g.search(cid, Q.account_totals(pdr))
    prev = {**agg(prev_rows), **(is_metrics(prev_rows[0]) if prev_rows else {})}
    d = deltas(cur, prev)
    d["search_is"] = pct_delta(cur.get("search_is"), prev.get("search_is"))

    # daily series
    daily = {x["segments.date"]: m(x) for x in g.search(cid, Q.account_daily(dr))}
    days = [(dr.start + timedelta(days=i)).isoformat() for i in range(dr.days)]
    cost_s = [daily.get(day, {}).get("cost") or 0 for day in days]
    conv_s = [daily.get(day, {}).get("conversions") or 0 for day in days]
    val_s = [daily.get(day, {}).get("conv_value") or 0 for day in days]
    kpi = tg["kpi"]
    roll = []
    for i in range(len(days)):
        c, v, n = sum(cost_s[max(0, i - 6):i + 1]), sum(conv_s[max(0, i - 6):i + 1]), sum(val_s[max(0, i - 6):i + 1])
        roll.append((c / v if v else None) if kpi == "cpa" else (n / c if c else None))

    # campaigns
    camp_cur, camp_prev = {}, {}
    for x in g.search(cid, Q.campaigns(dr)):
        camp_cur[str(x["campaign.id"])] = {"name": x["campaign.name"], "channel": x["campaign.advertising_channel_type"], **m(x)}
    for x in g.search(cid, Q.campaigns(pdr)):
        camp_prev[str(x["campaign.id"])] = m(x)
    camps = sorted(camp_cur.items(), key=lambda kv: -(kv[1]["cost"] or 0))
    channels: dict[str, list[dict[str, Any]]] = {}
    for x in g.search(cid, Q.campaigns(dr)):
        channels.setdefault(x["campaign.advertising_channel_type"], []).append(x)

    # wins / losses
    def eff(k: dict[str, Any]) -> float | None:
        return k.get("cpa") if kpi == "cpa" else k.get("roas")

    target_v = tg.get("target_cpa") if kpi == "cpa" else tg.get("target_roas")
    wins, losses = [], []
    for cid_, k in camps:
        p = camp_prev.get(cid_)
        dc = pct_delta(k["conversions"], p["conversions"]) if p else None
        e, pe = eff(k), eff(p) if p else None
        de = pct_delta(e, pe)
        good = (e is not None and target_v and ((kpi == "cpa" and e <= float(target_v)) or (kpi == "roas" and e >= float(target_v))))
        line = f"**{k['name']}** — {t['conversions'].lower()} {_fmt(k['conversions'])} ({_delta(dc)}), {kpi.upper()} {_fmt(e)} ({_delta(de)})"
        if (dc or 0) >= 15 or (de is not None and ((kpi == "cpa" and de <= -10) or (kpi == "roas" and de >= 10))) or good:
            wins.append((k["cost"] or 0, line))
        elif (k["cost"] or 0) > 0 and ((dc or 0) <= -15 or (de is not None and ((kpi == "cpa" and de >= 15) or (kpi == "roas" and de <= -15)))
                                      or (target_v and e is not None and not good)):
            losses.append((k["cost"] or 0, line))
    wins = [x for _, x in sorted(wins, reverse=True)[:5]]
    losses = [x for _, x in sorted(losses, reverse=True)[:5]]

    # actions from the audit log + decisions
    start_iso, end_iso = dr.start.isoformat(), (dr.end + timedelta(days=1)).isoformat()
    with db.connect() as conn:
        acts = [dict(x) for x in conn.execute(
            "SELECT p.id, p.kind, p.impact, p.reason, p.applied_at FROM proposal p WHERE p.customer_id=? AND "
            "p.status IN ('applied','partially_applied','rolled_back') AND p.applied_at>=? AND p.applied_at<? ORDER BY p.applied_at",
            (cid, start_iso, end_iso))]
        decs = [dict(x) for x in conn.execute(
            "SELECT ts, decision, rationale FROM decision WHERE customer_id=? AND ts>=? AND ts<? ORDER BY ts", (cid, start_iso, end_iso))]
    actions = [f"{a['applied_at'][:10]} — {(db.loads(a['impact']) or {}).get('summary') or a['kind']}" + (f" ({a['reason']})" if a.get("reason") else "")
               for a in acts]
    actions += [f"{x['ts'][:10]} — {x['decision']}" + (f": {x['rationale']}" if x.get("rationale") else "") for x in decs]

    next_steps: list[str] = []
    if include_next_steps:
        try:
            from tools.waste import analyze_wasted_spend
            w = analyze_wasted_spend.raw(cid, date_range=period, dimensions=["search_terms", "placements", "products"], limit=5)
            if w.get("rows"):
                tot = sum(x["wasted"] or 0 for x in w["rows"])
                next_steps.append(f"{t['waste']}: {_fmt(r(tot))} {info['currency']} — " + "; ".join(f"{x['entity']} ({_fmt(x['wasted'])})" for x in w["rows"][:3]))
        except Exception:  # noqa: BLE001 - next steps are best effort
            pass
        if tg.get("monthly_budget"):
            try:
                from tools.pacing import pace_account
                pc = pace_account(cid)
                next_steps.append(f"{t['pacing']}: {pc.get('pacing_pct')}% ({pc.get('status')})")
            except Exception:  # noqa: BLE001
                pass
        next_steps += [f"{t['losses']}: {re.sub(r'[*]', '', x.split(' — ')[0])}" for x in losses[:2]]

    # files
    name = accounts.alias_for(cid) or _slug(info.get("name") or cid)
    folder = Path(out_dir) if out_dir else get_settings().reports_dir / name / f"{dr.start.isoformat()}_{dr.end.isoformat()}_{lang}"
    folder.mkdir(parents=True, exist_ok=True)
    cur_code = info["currency"]
    ch1 = charts.daily_bars(days, cost_s, f"{t['daily_cost']}, {cur_code}", folder / "daily_cost.png")
    ch2 = charts.daily_line(days, conv_s, t["daily_conv"], folder / "daily_conversions.png")
    ch3 = charts.daily_line(days, roll, t["rolling_cpa"] if kpi == "cpa" else t["rolling_roas"], folder / "efficiency.png",
                            float(target_v) if target_v else None, t["target_line"])
    top = camps[:10]
    ch4 = charts.campaign_bars([k["name"] for _, k in top], [k["cost"] or 0 for _, k in top],
                               [f"{kpi.upper()} {_fmt(eff(k))}" for _, k in top], f"{t['by_campaign']}, {cur_code}", folder / "campaigns.png") if top else None
    chart_files = [c for c in (ch1, ch2, ch3, ch4) if c]

    kpi_rows = []
    for key in KPI_ROWS:
        tv = tg.get("target_cpa") if key == "cpa" else tg.get("target_roas") if key == "roas" else None
        kpi_rows.append([t[key], _fmt(cur.get(key), key), _fmt(prev.get(key), key), _delta(d.get(key)), _fmt(tv) if tv else ""])
    camp_rows = [[k["name"], k["channel"], _fmt(k["cost"]), _fmt(k["conversions"]), _fmt(k["cpa"]), _fmt(k["roas"]),
                  _delta(pct_delta(k["cost"], camp_prev.get(c, {}).get("cost"))),
                  _delta(pct_delta(k["conversions"], camp_prev.get(c, {}).get("conversions")))] for c, k in camps[:25]]
    chan_rows = [[ch, _fmt(agg(v)["cost"]), _fmt(agg(v)["conversions"]), _fmt(agg(v)["cpa"]), _fmt(agg(v)["roas"])]
                 for ch, v in sorted(channels.items(), key=lambda kv: -(agg(kv[1])["cost"] or 0))]

    eff_line = ""
    if target_v and eff(cur) is not None:
        ok = (kpi == "cpa" and eff(cur) <= float(target_v)) or (kpi == "roas" and eff(cur) >= float(target_v))
        eff_line = f" {kpi.upper()} {_fmt(eff(cur))} vs {t['target'].lower()} {_fmt(float(target_v))} — {t['on_target'] if ok else t['off_target']}."
    summary = (f"{t['cost']}: {_fmt(cur['cost'])} {cur_code} ({_delta(d.get('cost'))}); {t['conversions'].lower()}: "
               f"{_fmt(cur['conversions'])} ({_delta(d.get('conversions'))}); {t['conv_value'].lower()}: {_fmt(cur['conv_value'])} "
               f"{cur_code} ({_delta(d.get('conv_value'))}).{eff_line}")

    head = account_header(cid, info, dr)
    title = f"{t['title']}: {info.get('name')}"
    period_line = f"{t['period']}: {dr.start} — {dr.end} ({t['compared']} {pdr.start} — {pdr.end})"
    md = [f"# {title}", "", period_line, "", f"## {t['summary']}", "", summary, "",
          f"## {t['kpis']}", "", _md_table([t["metric"], t["value"], t["prev"], t["delta"], t["target"]], kpi_rows), "",
          f"## {t['charts']}", ""] + [f"![{Path(c).stem}]({Path(c).name})" for c in chart_files] + [
          "", f"## {t['channels']}", "", _md_table([t["channel"], t["cost"], t["conversions"], "CPA", "ROAS"], chan_rows), "",
          f"## {t['campaigns']}", "", _md_table([t["campaign"], t["channel"], t["cost"], t["conversions"], "CPA", "ROAS",
                                                f"Δ {t['cost'].lower()}", f"Δ {t['conversions'].lower()}"], camp_rows), "",
          f"## {t['wins']}", ""] + [f"- {x}" for x in wins or ["—"]] + ["", f"## {t['losses']}", ""] + \
         [f"- {x}" for x in losses or ["—"]] + ["", f"## {t['actions']}", ""] + [f"- {x}" for x in actions or [t["no_actions"]]] + \
         ["", f"## {t['next']}", ""] + [f"- {x}" for x in next_steps or ["—"]] + \
         ["", f"_{t['generated']}: {datetime.now().strftime('%Y-%m-%d %H:%M')} · adsops-mcp_"]
    md_text = "\n".join(md)
    md_path = folder / "report.md"
    md_path.write_text(md_text, encoding="utf-8")
    html_path = folder / "report.html"
    html_path.write_text(_html(md_text, folder, title), encoding="utf-8")
    xlsx_path = _xlsx(folder / "report.xlsx", t, kpi_rows, camp_rows, chan_rows, days, cost_s, conv_s, val_s) if xlsx else None

    return {**head, "files": {"markdown": str(md_path), "html": str(html_path), "xlsx": str(xlsx_path) if xlsx_path else None,
                              "charts": chart_files},
            "summary": summary, "kpis": cur, "previous": prev, "delta_pct": d, "wins": wins, "losses": losses,
            "actions_taken": actions, "next_steps": next_steps, "lang": lang}


def _inline(s: str) -> str:
    s = html.escape(s)
    s = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", s)
    return re.sub(r"_(.+?)_$", r"<em>\1</em>", s)


def _html(md: str, folder: Path, title: str) -> str:
    """Tiny Markdown subset -> standalone HTML with base64-embedded charts (one file to send to the client)."""
    body, table, in_list = [], [], False

    def flush_table():
        if not table:
            return
        rows = [[c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", row.strip()[1:-1])]
                for row in table if not re.match(r"^\|[-|]+\|$", row)]
        out = ["<div class='tw'><table><thead><tr>" + "".join(f"<th>{_inline(c)}</th>" for c in rows[0]) + "</tr></thead><tbody>"]
        out += ["<tr>" + "".join(f"<td>{_inline(c)}</td>" for c in r_) + "</tr>" for r_ in rows[1:]]
        body.append("".join(out) + "</tbody></table></div>")
        table.clear()

    for line in md.splitlines():
        if line.startswith("|"):
            table.append(line)
            continue
        flush_table()
        if in_list and not line.startswith("- "):
            body.append("</ul>")
            in_list = False
        if line.startswith("# "):
            body.append(f"<h1>{_inline(line[2:])}</h1>")
        elif line.startswith("## "):
            body.append(f"<h2>{_inline(line[3:])}</h2>")
        elif line.startswith("!["):
            src = re.search(r"\((.+?)\)", line).group(1)
            data = base64.b64encode((folder / src).read_bytes()).decode()
            body.append(f"<img alt='{html.escape(src)}' src='data:image/png;base64,{data}'>")
        elif line.startswith("- "):
            if not in_list:
                body.append("<ul>")
                in_list = True
            body.append(f"<li>{_inline(line[2:])}</li>")
        elif line.strip():
            body.append(f"<p>{_inline(line)}</p>")
    flush_table()
    if in_list:
        body.append("</ul>")
    css = """:root{--bg:#f9f9f7;--surface:#fcfcfb;--ink:#0b0b0b;--ink2:#52514e;--line:#e4e3df;--accent:#2a78d6}
@media (prefers-color-scheme: dark){:root{--bg:#0d0d0d;--surface:#1a1a19;--ink:#f0efec;--ink2:#c3c2b7;--line:#383835;--accent:#3987e5}
img{filter:brightness(.92)}}
body{margin:0;background:var(--bg);color:var(--ink);font:15px/1.55 -apple-system,Segoe UI,Roboto,Arial,sans-serif}
main{max-width:920px;margin:0 auto;padding:32px 16px}h1{font-size:26px;margin:0 0 4px}h2{font-size:18px;margin:32px 0 10px;
border-bottom:1px solid var(--line);padding-bottom:6px}p,li{color:var(--ink)}p:first-of-type{color:var(--ink2)}
img{width:100%;height:auto;border-radius:8px;background:#fcfcfb;margin:8px 0}.tw{overflow-x:auto}
table{border-collapse:collapse;width:100%;font-size:13px;background:var(--surface)}th,td{padding:7px 10px;
border-bottom:1px solid var(--line);text-align:right;white-space:nowrap}th:first-child,td:first-child{text-align:left;
white-space:normal}th{color:var(--ink2);font-weight:600}"""
    return (f"<!doctype html><html><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
            f"<title>{html.escape(title)}</title><style>{css}</style></head><body><main>{''.join(body)}</main></body></html>")


def _xlsx(path: Path, t, kpi_rows, camp_rows, chan_rows, days, cost_s, conv_s, val_s) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    sheets = [(t["kpis"], [t["metric"], t["value"], t["prev"], t["delta"], t["target"]], kpi_rows),
              (t["campaigns"], [t["campaign"], t["channel"], t["cost"], t["conversions"], "CPA", "ROAS", "Δ cost", "Δ conv"], camp_rows),
              (t["channels"], [t["channel"], t["cost"], t["conversions"], "CPA", "ROAS"], chan_rows),
              ("Daily", ["date", t["cost"], t["conversions"], t["conv_value"]],
               [[d_, c, v, n] for d_, c, v, n in zip(days, cost_s, conv_s, val_s)])]
    for i, (name, head, rows) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = name[:31]
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
        for row in rows:
            ws.append([_num(x) for x in row])
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(10, min(45, max(len(str(c.value or "")) for c in col) + 2))
    wb.save(path)
    return path


def _num(v: Any) -> Any:
    if isinstance(v, str):
        s = v.replace(" ", "").replace("%", "").replace("+", "")
        try:
            return float(s)
        except ValueError:
            return v
    return v
